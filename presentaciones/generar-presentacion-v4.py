"""Generador de la presentación v4 del Seminario IIE-3T con python-pptx.

Incluye ilustraciones científicas en español, conectores vectoriales (arcos sólidos y punteados/discontinuos)
en todos los grafos, eliminación rigurosa de puntos finales en títulos y lemas, y cajas ajustadas
para evitar cualquier desbordamiento de texto.
"""
from pathlib import Path
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "output"
OUTPUT_DIR.mkdir(exist_ok=True)
PPTX_PATH = OUTPUT_DIR / "Introduccion-iie3t-redes-bayesianas-v4.pptx"

# Colores institucionales del seminario
BG_COLOR = RGBColor(0xFA, 0xFA, 0xF7)       # Fondo cálido claro
INK_COLOR = RGBColor(0x18, 0x34, 0x3A)      # Texto principal oscuro
MUTED_COLOR = RGBColor(0x52, 0x65, 0x6A)    # Texto secundario
CONTEXT_COLOR = RGBColor(0x08, 0x7E, 0x86)  # Turquesa (Contexto X)
OBS_COLOR = RGBColor(0xAA, 0x67, 0x17)      # Ocre (Detección Y)
LATENT_COLOR = RGBColor(0x73, 0x52, 0x8C)   # Violeta (Condición Z)
GRAY_BG = RGBColor(0xDD, 0xE5, 0xE6)        # Gris observado
MATH_BG = RGBColor(0xEE, 0xF4, 0xF5)        # Fondo para cajas matemáticas
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT_GREEN = RGBColor(0x2E, 0x7D, 0x32)
LINE_COLOR = RGBColor(0x52, 0x65, 0x6A)

FONT_MAIN = "Arial"
FONT_MATH = "Cambria Math"


def add_arrow(slide, x1, y1, x2, y2, dashed=False, color=LINE_COLOR, width=Pt(2.5)):
    """Crea un conector recto con punta de flecha en el extremo de destino.
    
    En DrawingML, los elementos hijos de <a:ln> deben seguir un orden estricto de esquema:
    1. <a:prstDash> (estilo de trazo punteado/discontinuo)
    2. <a:tailEnd> (punta de flecha)
    """
    cxn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, x1, y1, x2, y2)
    cxn.line.color.rgb = color
    cxn.line.width = width
    spPr = cxn._element.spPr
    ln = spPr.find('{http://schemas.openxmlformats.org/drawingml/2006/main}ln')

    if dashed:
        dash = OxmlElement('a:prstDash')
        dash.set('val', 'dash')
        ln.append(dash)

    # En DrawingML, tailEnd es la punta de flecha en el punto final (x2, y2)
    tail = OxmlElement('a:tailEnd')
    tail.set('type', 'triangle')
    tail.set('w', 'med')
    tail.set('len', 'med')
    ln.append(tail)

    return cxn


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_base_slide(title_text, slide_num, notes_text=""):
        slide = prs.slides.add_slide(blank_layout)
        # Background
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

        # Title (sin punto final)
        clean_title = title_text.rstrip(".") if title_text else ""
        if clean_title:
            tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.9))
            tf = tb.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            p = tf.paragraphs[0]
            p.text = clean_title
            p.font.name = FONT_MAIN
            p.font.size = Pt(28)
            p.font.bold = True
            p.font.color.rgb = INK_COLOR

        # Slide Number
        tb_num = slide.shapes.add_textbox(Inches(12.0), Inches(6.9), Inches(0.8), Inches(0.4))
        tf_num = tb_num.text_frame
        p_num = tf_num.paragraphs[0]
        p_num.text = str(slide_num).zfill(2)
        p_num.font.name = FONT_MAIN
        p_num.font.size = Pt(14)
        p_num.font.color.rgb = MUTED_COLOR
        p_num.alignment = PP_ALIGN.RIGHT

        # Speaker notes
        if notes_text:
            slide.notes_slide.notes_text_frame.text = notes_text

        return slide

    # -------------------------------------------------------------
    # SLIDE 1: Portada
    # -------------------------------------------------------------
    s1 = add_base_slide("", 1,
        "Notas del facilitador: Bienvenida al taller del Seminario IIE-3T. "
        "En esta sesión introductoria tenderemos un puente conceptual desde los fundamentos de salud e integridad "
        "ecosistémica (enfoque Ecosalud / Margulis / Equihua et al.) y el pensamiento sistémico (Forrester, Meadows) "
        "hacia el modelado probabilístico con Redes Bayesianas y el modelo en tres capas del equipo.")

    # Header Box (sin puntos en títulos ni subtítulos)
    tb1 = s1.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(1.6))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p1_1 = tf1.paragraphs[0]
    p1_1.text = "De las señales del ecosistema a un diagnóstico que reconoce las incertidumbres"
    p1_1.font.name = FONT_MAIN
    p1_1.font.size = Pt(28)
    p1_1.font.bold = True
    p1_1.font.color.rgb = INK_COLOR

    p1_2 = tf1.add_paragraph()
    p1_2.text = "Integridad Ecosistémica, Modelo de Tres Capas (IIE-3T) y Redes Bayesianas"
    p1_2.font.name = FONT_MAIN
    p1_2.font.size = Pt(18)
    p1_2.font.color.rgb = CONTEXT_COLOR
    p1_2.space_before = Pt(6)

    # Hero Image en español
    img_portada = ROOT / "presentaciones/img/portada_integridad_redes_es.jpg"
    if not img_portada.exists():
        img_portada = ROOT / "presentaciones/img/portada_integridad_redes.jpg"
    if img_portada.exists():
        s1.shapes.add_picture(str(img_portada), Inches(0.8), Inches(2.3), Inches(11.733), Inches(4.5))

    # -------------------------------------------------------------
    # SLIDE 2: La integridad como salud ecosistémica
    # -------------------------------------------------------------
    s2 = add_base_slide("La integridad como salud del ecosistema", 2,
        "Fuentes: Colección Ecosalud (UNAM), Margulis (1997), Equihua et al. "
        "Idea central: La integridad no es una colección de partes ni una superficie verde. "
        "Es la capacidad de autorregulación y autopoiesis del tejido vivo. "
        "La salud es una variable latente: no se mide directamente con un termómetro único, "
        "sino que se diagnostica mediante múltiples señales observables en contexto.")

    # Left Column: Analogías y Concepto
    tb2_left = s2.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(5.6), Inches(5.2))
    tf2_l = tb2_left.text_frame
    tf2_l.word_wrap = True

    p = tf2_l.paragraphs[0]
    p.text = "Más allá de una apariencia verde"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR

    p = tf2_l.add_paragraph()
    p.text = "• Un campo de golf es verde y un atleta hipertrofiado luce fuerte, pero ninguno garantiza salud funcional autorregulada."
    p.font.size = Pt(15.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(8)

    p = tf2_l.add_paragraph()
    p.text = "• La integridad ecosistémica describe la condición viva y autopoiética de la red ecológica completa."
    p.font.size = Pt(15.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(8)

    p = tf2_l.add_paragraph()
    p.text = "• No es un inventario estático de especies, sino la persistencia de procesos ecológicos coordinados."
    p.font.size = Pt(15.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(8)

    # Right Column: 3 Pilares
    box_w = Inches(5.6)
    box_h = Inches(1.5)
    for idx, (title_box, desc_box, col) in enumerate([
        ("1. Autopoiesis y metabolismo", "Capacidad intrínseca de regenerar componentes, fijar energía y reciclar nutrientes.", CONTEXT_COLOR),
        ("2. Autorregulación organísmica", "Amortiguamiento de perturbaciones y estabilidad funcional dinámica.", OBS_COLOR),
        ("3. Condición como estado latente", "La salud no se observa directamente: se diagnostica evaluando señales en contexto.", LATENT_COLOR),
    ]):
        top_y = Inches(1.45 + idx * 1.65)
        rect = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), top_y, box_w, box_h)
        rect.fill.solid()
        rect.fill.fore_color.rgb = WHITE
        rect.line.color.rgb = col
        rect.line.width = Pt(2)

        tf_box = rect.text_frame
        tf_box.word_wrap = True
        tf_box.margin_left = tf_box.margin_right = Inches(0.2)
        tf_box.margin_top = Inches(0.12)

        pb1 = tf_box.paragraphs[0]
        pb1.text = title_box.rstrip(".")
        pb1.font.name = FONT_MAIN
        pb1.font.size = Pt(16.5)
        pb1.font.bold = True
        pb1.font.color.rgb = col

        pb2 = tf_box.add_paragraph()
        pb2.text = desc_box
        pb2.font.name = FONT_MAIN
        pb2.font.size = Pt(13.5)
        pb2.font.color.rgb = INK_COLOR
        pb2.space_before = Pt(3)

    # -------------------------------------------------------------
    # SLIDE 3: Pensamiento sistémico y el microcosmos viviente
    # -------------------------------------------------------------
    s3 = add_base_slide("Pensamiento sistémico y el microcosmos viviente", 3,
        "Fuentes: Forrester (1961), Sterman (2000), Meadows (2008), Ecosalud. "
        "Analogía de la pecera: En un vivario acuático, el equilibrio hídrico no es una pieza aislada. "
        "La claridad del agua y la salud biótica dependen de bucles de fotosíntesis, mineralización de amonio y biofiltración. "
        "El pensamiento sistémico nos invita a modelar las relaciones de interdependencia antes de calcular índices.")

    tb3_left = s3.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.2))
    tf3 = tb3_left.text_frame
    tf3.word_wrap = True

    p = tf3.paragraphs[0]
    p.text = "La analogía de la pecera (Vivario)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR

    p = tf3.add_paragraph()
    p.text = "• En un acuario plantado, el agua limpia y los peces sanos no se logran agregando productos químicos aislados."
    p.font.size = Pt(15)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf3.add_paragraph()
    p.text = "• La salud es el equilibrio dinámico entre fotosíntesis, bacterias nitrificantes, sustrato y cargas orgánicas."
    p.font.size = Pt(15)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf3.add_paragraph()
    p.text = "El todo y sus dependencias (Forrester / Meadows)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(12)

    p = tf3.add_paragraph()
    p.text = "• Todo está interconectado, pero la estructura de influencia no es uniforme: existen vías directas y mediadas."
    p.font.size = Pt(15)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    p = tf3.add_paragraph()
    p.text = "• Para diagnosticar el sistema necesitamos formalizar explícitamente estas vías de interacción."
    p.font.size = Pt(15)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # Right Image: Microcosmos Pecera en español
    img_pecera = ROOT / "presentaciones/img/sistemico_pecera_microcosmos_es.jpg"
    if not img_pecera.exists():
        img_pecera = ROOT / "presentaciones/img/sistemico_pecera_microcosmos.jpg"
    if img_pecera.exists():
        s3.shapes.add_picture(str(img_pecera), Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.2))

    # -------------------------------------------------------------
    # SLIDE 4: Tres capas analíticas: IIE-3T
    # -------------------------------------------------------------
    s4 = add_base_slide("Tres capas analíticas: Modelo IIE-3T", 4,
        "Fuente: iie-teoria/iie-teoria.qmd, §§3–5. "
        "Las capas Contexto (X), Detección (Y) y Condición latente (Z) constituyen el núcleo analítico. "
        "Permiten no confundir las condiciones de partida con las observaciones o el estado biótico subyacente.")

    card_w = Inches(3.6)
    card_h = Inches(4.85)
    cards_data = [
        ("Contexto (X)", CONTEXT_COLOR, [
            ("Lo que el ecosistema recibe", True),
            ("Entorno físico-químico y biogeográfico de fondo.", False),
            ("• Macroclima y precipitación", False),
            ("• Relieve y topografía", False),
            ("• Tipo de suelo y geología", False),
            ("• Zona de vida biogeográfica", False),
        ]),
        ("Detección (Y)", OBS_COLOR, [
            ("Las señales observables", True),
            ("Mediciones instrumentales y datos de monitoreo.", False),
            ("• Índices espectrales (NDVI/EVI)", False),
            ("• Cobertura vegetal y biomasa", False),
            ("• Inventarios y muestreos de campo", False),
            ("• Registros de sensores remotos", False),
        ]),
        ("Condición (Z)", LATENT_COLOR, [
            ("El estado biótico latente", True),
            ("La integridad y funcionalidad ecológica real.", False),
            ("• No observable directamente", False),
            ("• Estado de salud sistémica", False),
            ("• Capacidad de autorregulación", False),
            ("• Objetivo del diagnóstico", False),
        ]),
    ]

    for idx, (title_c, col, items) in enumerate(cards_data):
        pos_x = Inches(0.8 + idx * 4.05)
        rect = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, pos_x, Inches(1.45), card_w, card_h)
        rect.fill.solid()
        rect.fill.fore_color.rgb = WHITE
        rect.line.color.rgb = col
        rect.line.width = Pt(2.5)

        tf = rect.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.18)

        p = tf.paragraphs[0]
        p.text = title_c.rstrip(".")
        p.font.name = FONT_MAIN
        p.font.size = Pt(21)
        p.font.bold = True
        p.font.color.rgb = col

        for text_item, is_bold in items:
            p = tf.add_paragraph()
            p.text = text_item.rstrip(".") if is_bold else text_item
            p.font.name = FONT_MAIN
            p.font.size = Pt(14 if not is_bold else 15)
            p.font.bold = is_bold
            p.font.color.rgb = INK_COLOR if not is_bold else col
            p.space_before = Pt(6 if is_bold else 3)

    # Footer note (sin punto final)
    tb4_foot = s4.shapes.add_textbox(Inches(0.8), Inches(6.45), Inches(11.7), Inches(0.4))
    p = tb4_foot.text_frame.paragraphs[0]
    p.text = "Distinción crucial: El contexto no es degradación; las señales observables no son el estado de integridad en sí mismo"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13.5)
    p.font.italic = True
    p.font.color.rgb = MUTED_COLOR

    # -------------------------------------------------------------
    # SLIDE 5: El grafo generativo mínimo (con arcos vectoriales completos)
    # -------------------------------------------------------------
    s5 = add_base_slide("El proceso generativo: cómo la naturaleza produce señales", 5,
        "Fuente: iie-teoria/iie-teoria.qmd, §§3–5 y 21. "
        "En la naturaleza, el contexto X y la condición biótica real Z determinan conjuntamente las observaciones Y. "
        "Nodos grises: observados; nodo blanco/violeta: latente. "
        "Acuerdo didáctico del equipo: existen argumentos conceptuales para omitir el arco X→Z, pero no una solución definitiva. "
        "La línea discontinua es una anotación editorial de incertidumbre estructural, no una flecha débil ni un nuevo tipo de dependencia probabilística. "
        "Para calcular, cada modelo debe incluir o excluir el arco de forma explícita. Las flechas continuas describen la estructura de trabajo y tampoco prueban causalidad. "
        "Mencionar esta incertidumbre brevemente en la introducción. Reservar el contraste entre modelos para una actividad posterior, "
        "tras estudiar separación-d, identificación de variables latentes y validación. No presentar un contraste no significativo como demostración de independencia.")

    # Left: Diagram of DAG
    # Context X (top-left): centro en (2.25, 2.55)
    node_x = s5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.5), Inches(1.8), Inches(1.5), Inches(1.5))
    node_x.fill.solid()
    node_x.fill.fore_color.rgb = GRAY_BG
    node_x.line.color.rgb = CONTEXT_COLOR
    node_x.line.width = Pt(3)
    tf_nx = node_x.text_frame
    p = tf_nx.paragraphs[0]
    p.text = "X\nContexto"
    p.font.name = FONT_MAIN
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Latent Z (bottom-left): centro en (2.25, 5.25)
    node_z = s5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.5), Inches(4.5), Inches(1.5), Inches(1.5))
    node_z.fill.solid()
    node_z.fill.fore_color.rgb = WHITE
    node_z.line.color.rgb = LATENT_COLOR
    node_z.line.width = Pt(3)
    tf_nz = node_z.text_frame
    p = tf_nz.paragraphs[0]
    p.text = "Z\nCondición\n(Latente)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Observations Y (middle-right): centro en (5.35, 3.90)
    node_y = s5.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.6), Inches(3.15), Inches(1.5), Inches(1.5))
    node_y.fill.solid()
    node_y.fill.fore_color.rgb = GRAY_BG
    node_y.line.color.rgb = OBS_COLOR
    node_y.line.width = Pt(3)
    tf_ny = node_y.text_frame
    p = tf_ny.paragraphs[0]
    p.text = "Y\nDetección\n(Señales)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = OBS_COLOR
    p.alignment = PP_ALIGN.CENTER

    # ARCOS VECTORIALES
    # 1. X -> Y (sólido)
    add_arrow(s5, Inches(2.9), Inches(2.3), Inches(4.7), Inches(3.4), dashed=False, color=CONTEXT_COLOR)
    # 2. Z -> Y (sólido)
    add_arrow(s5, Inches(2.9), Inches(4.9), Inches(4.7), Inches(4.2), dashed=False, color=LATENT_COLOR)
    # 3. X -> Z (discontinuo/punteado, en evaluación)
    add_arrow(s5, Inches(2.25), Inches(3.3), Inches(2.25), Inches(4.5), dashed=True, color=LINE_COLOR)

    # Right Box: Explicación y Leyenda
    tb5_right = s5.shapes.add_textbox(Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.2))
    tf5_r = tb5_right.text_frame
    tf5_r.word_wrap = True

    p = tf5_r.paragraphs[0]
    p.text = "Dirección generativa natural"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR

    p = tf5_r.add_paragraph()
    p.text = "• X → Y (Arco directo): El clima y el suelo modulan el tipo y magnitud de las señales observables."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf5_r.add_paragraph()
    p.text = "• Z → Y (Arco de emisión): El estado real de salud ecológica genera los síntomas observables."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf5_r.add_paragraph()
    p.text = "• X ⇢ Z (En evaluación): ¿El contexto determina directamente el potencial de salud o fija la referencia basal? Es una hipótesis abierta."
    p.font.size = Pt(14.5)
    p.font.color.rgb = MUTED_COLOR
    p.space_before = Pt(6)

    p = tf5_r.add_paragraph()
    p.text = "Convención visual del grafo"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(10)

    p = tf5_r.add_paragraph()
    p.text = "⚪ Gris = Variable observada   |   ⚪ Blanco = Variable latente\n⟶ Sólida = Dependencia fijada  |   ⇢ Discontinua = Hipótesis en evaluación"
    p.font.size = Pt(12.5)
    p.font.color.rgb = MUTED_COLOR
    p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 6: El DAG como mapa de la probabilidad conjunta (Ecuaciones estilizadas)
    # -------------------------------------------------------------
    s6 = add_base_slide("El DAG: Factorización de la Probabilidad Conjunta", 6,
        "Fundamento matemático: Pearl (1988), Lauritzen (1996). "
        "Un Grafo Acíclico Dirigido no es solo un diagrama conceptual: es la regla matemática exacta "
        "para descomponer una distribución conjunta masiva en el producto de distribuciones condicionales locales.")

    # Top text
    tb6_top = s6.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.1))
    tf6_t = tb6_top.text_frame
    tf6_t.word_wrap = True
    p = tf6_t.paragraphs[0]
    p.text = "Descomposición de la complejidad del ecosistema"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR

    p = tf6_t.add_paragraph()
    p.text = "Modelar un ecosistema con decenas de variables de forma conjunta sería intratable. La regla de factorización de un DAG simplifica la distribución conjunta global en factores locales condicionados únicamente a los padres:"
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # Box 1: Regla General de Factorización
    f_box1 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.5), Inches(10.3), Inches(1.35))
    f_box1.fill.solid()
    f_box1.fill.fore_color.rgb = MATH_BG
    f_box1.line.color.rgb = CONTEXT_COLOR
    f_box1.line.width = Pt(1.5)

    tf_fb1 = f_box1.text_frame
    tf_fb1.word_wrap = True
    tf_fb1.margin_top = Inches(0.12)
    p = tf_fb1.paragraphs[0]
    p.text = "Regla General de Factorización en un DAG"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR
    p.alignment = PP_ALIGN.CENTER

    p = tf_fb1.add_paragraph()
    p.text = "P( X₁, X₂, ..., Xₙ )  =  ∏ᵢ₌₁ⁿ  P( Xᵢ | Padres(Xᵢ) )"
    p.font.name = FONT_MATH
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(3)

    # Box 2: Factorización para el Modelo IIE-3T
    f_box2 = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(4.0), Inches(10.3), Inches(1.35))
    f_box2.fill.solid()
    f_box2.fill.fore_color.rgb = MATH_BG
    f_box2.line.color.rgb = LATENT_COLOR
    f_box2.line.width = Pt(1.5)

    tf_fb2 = f_box2.text_frame
    tf_fb2.word_wrap = True
    tf_fb2.margin_top = Inches(0.12)
    p = tf_fb2.paragraphs[0]
    p.text = "Factorización para el Modelo IIE-3T Mínimo (X, Y, Z)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR
    p.alignment = PP_ALIGN.CENTER

    p = tf_fb2.add_paragraph()
    p.text = "P( X, Y, Z )  =  P( X )  ·  P( Z | X )  ·  P( Y | X, Z )"
    p.font.name = FONT_MATH
    p.font.size = Pt(19)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR
    p.alignment = PP_ALIGN.CENTER
    p.space_before = Pt(3)

    # Bottom explanatory bullets
    tb6_bot = s6.shapes.add_textbox(Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.5))
    tf6_b = tb6_bot.text_frame
    tf6_b.word_wrap = True
    p = tf6_b.paragraphs[0]
    p.text = "• Cada nodo solo requiere una tabla local de probabilidades condicionada a sus antecedentes directos."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p = tf6_b.add_paragraph()
    p.text = "• Independencia condicional (separación-d): El grafo define formalmente qué variables quedan bloqueadas de influencia al conocer cierta evidencia."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 7: Tablas de Probabilidad Condicional (CPT)
    # -------------------------------------------------------------
    s7 = add_base_slide("El componente numérico: Tablas Condicionales (CPT)", 7,
        "Fuente: Cafe-blog/posts/redes-bayesianas-basico/index.qmd y docs/REVISION-CONCEPTUAL-CLIMA.md. "
        "Cada nodo del DAG posee una Tabla de Probabilidad Condicional (CPT). "
        "Aquí se muestra el caso climático adaptado: la distribución de Cobertura vegetal (Cob) condicionada a Árboles (A) y Cambio Climático (C).")

    tb7_top = s7.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(0.9))
    tf7 = tb7_top.text_frame
    tf7.word_wrap = True
    p = tf7.paragraphs[0]
    p.text = "Ejemplo: P( Cobertura | Árboles, Cambio Climático )"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = OBS_COLOR

    # Native Table
    rows = 5
    cols = 4
    left = Inches(0.8)
    top = Inches(2.3)
    width = Inches(11.7)
    height = Inches(2.8)

    table_shape = s7.shapes.add_table(rows, cols, left, top, width, height)
    tbl = table_shape.table
    tbl.columns[0].width = Inches(2.6)
    tbl.columns[1].width = Inches(2.6)
    tbl.columns[2].width = Inches(3.25)
    tbl.columns[3].width = Inches(3.25)

    headers = ["Árboles (A)", "Cambio Climático (C)", "P( Cobertura = Densa )", "P( Cobertura = Rala )"]
    data = [
        ["abundante", "moderado", "0.90", "0.10"],
        ["abundante", "severo", "0.75", "0.25"],
        ["escaso", "moderado", "0.30", "0.70"],
        ["escaso", "severo", "0.05", "0.95"],
    ]

    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        cell.fill.solid()
        cell.fill.fore_color.rgb = INK_COLOR
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.name = FONT_MAIN
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.alignment = PP_ALIGN.CENTER

    for r_idx, row_data in enumerate(data):
        for c_idx, val in enumerate(row_data):
            cell = tbl.cell(r_idx + 1, c_idx)
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r_idx % 2 == 0 else GRAY_BG
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = FONT_MAIN
            p.font.size = Pt(15)
            p.font.color.rgb = INK_COLOR
            p.alignment = PP_ALIGN.CENTER

    tb7_foot = s7.shapes.add_textbox(Inches(0.8), Inches(5.45), Inches(11.7), Inches(1.3))
    tf7_f = tb7_foot.text_frame
    tf7_f.word_wrap = True
    p = tf7_f.paragraphs[0]
    p.text = "• Cada fila suma exactamente 1.0 (conservación de la probabilidad)."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p = tf7_f.add_paragraph()
    p.text = "• Las CPTs pueden derivarse de frecuencias empíricas, juicio experto calibrado o ecuaciones de proceso."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 8: Inferencia diagnóstica (Ecuación estructurada de Bayes)
    # -------------------------------------------------------------
    s8 = add_base_slide("Inferencia diagnóstica: de las señales a la salud", 8,
        "Fuente: iie-teoria §5.1 y §21. "
        "El ecólogo observa síntomas Y en un contexto X, e invierte probabilísticamente el proceso con la Regla de Bayes "
        "para inferir la distribución posterior de la condición latente Z. "
        "Una probabilidad posterior de integridad de 0.72 expresa certeza diagnóstica, no 72% de masa física conservada.")

    tb8_top = s8.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.0))
    tf8_t = tb8_top.text_frame
    tf8_t.word_wrap = True
    p = tf8_t.paragraphs[0]
    p.text = "El razonamiento clínico en ecología"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR

    p = tf8_t.add_paragraph()
    p.text = "La naturaleza opera de la salud a los síntomas (Z → Y). El ecólogo trabaja en sentido inverso: a partir de las señales observables Y y el contexto X, diagnostica la condición latente Z mediante el Teorema de Bayes:"
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(3)

    # Elegant Bayes Equation Card
    bayes_card = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(2.45), Inches(10.3), Inches(1.8))
    bayes_card.fill.solid()
    bayes_card.fill.fore_color.rgb = MATH_BG
    bayes_card.line.color.rgb = LATENT_COLOR
    bayes_card.line.width = Pt(1.5)

    tf_bc = bayes_card.text_frame
    tf_bc.word_wrap = True
    tf_bc.margin_top = Inches(0.12)
    p = tf_bc.paragraphs[0]
    p.text = "Inversión Probabilística Diagnóstica"
    p.font.name = FONT_MAIN
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Fraction Display with clear numerator / denominator
    p = tf_bc.add_paragraph()
    p.text = "                                       P( Y = y | Z, X )  ·  P( Z | X )"
    p.font.name = FONT_MATH
    p.font.size = Pt(16.5)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR
    p.space_before = Pt(2)

    p = tf_bc.add_paragraph()
    p.text = "P( Z | Y = y, X = x )  =  ──────────────────────────────────────────"
    p.font.name = FONT_MATH
    p.font.size = Pt(16.5)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR

    p = tf_bc.add_paragraph()
    p.text = "                                                   P( Y = y | X )"
    p.font.name = FONT_MATH
    p.font.size = Pt(16.5)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR

    # Bottom bullets
    tb8_bot = s8.shapes.add_textbox(Inches(0.8), Inches(4.45), Inches(11.7), Inches(2.5))
    tf8_b = tb8_bot.text_frame
    tf8_b.word_wrap = True

    p = tf8_b.paragraphs[0]
    p.text = "1. Distribución a priori P(Z | X): Conocimiento basal de la condición según el contexto ecológico."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR

    p = tf8_b.add_paragraph()
    p.text = "2. Verosimilitud P(Y | Z, X): Qué tan compatibles son los datos observados bajo cada nivel de salud."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    p = tf8_b.add_paragraph()
    p.text = "3. Distribución a posteriori P(Z | Y, X): El diagnóstico que cuantifica la certidumbre sin forzar un valor determinista falso."
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 9: Notación compacta de plates (con arcos vectoriales)
    # -------------------------------------------------------------
    s9 = add_base_slide("Múltiples sitios: Notación compacta de Plates", 9,
        "Referencia de notación: Blei & Lafferty (2009). "
        "El rectángulo (plate) denota repetición para i = 1,...,N sitios de muestreo o píxeles. "
        "θ representa los parámetros globales del modelo que se comparten entre todos los sitios. "
        "El plate expresa repetición muestral, no relaciones causales ni dependencia espacial automática. "
        "Acuerdo didáctico del equipo: Existen argumentos conceptuales para omitir el arco X→Z, pero no una solución definitiva. "
        "La línea discontinua es una anotación editorial de incertidumbre estructural, no una flecha débil ni un nuevo tipo de dependencia probabilística. "
        "Para calcular, cada modelo formal debe incluir o excluir el arco de forma explícita.")

    # Plate Box: left=1.0, top=1.8, width=5.2, height=4.7
    plate = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.8), Inches(5.2), Inches(4.7))
    plate.fill.background()
    plate.line.color.rgb = MUTED_COLOR
    plate.line.width = Pt(2)

    tb_plate_lbl = s9.shapes.add_textbox(Inches(1.2), Inches(5.95), Inches(4.8), Inches(0.4))
    p = tb_plate_lbl.text_frame.paragraphs[0]
    p.text = "Sitios de muestreo   i = 1, ..., N"
    p.font.name = FONT_MAIN
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR

    # Nodes inside plate
    # Xi: centro en (2.2, 2.7)
    nx = s9.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.65), Inches(2.15), Inches(1.1), Inches(1.1))
    nx.fill.solid()
    nx.fill.fore_color.rgb = GRAY_BG
    nx.line.color.rgb = CONTEXT_COLOR
    nx.line.width = Pt(2.5)
    p = nx.text_frame.paragraphs[0]
    p.text = "Xᵢ"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Zi: centro en (2.2, 4.8)
    nz = s9.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.65), Inches(4.25), Inches(1.1), Inches(1.1))
    nz.fill.solid()
    nz.fill.fore_color.rgb = WHITE
    nz.line.color.rgb = LATENT_COLOR
    nz.line.width = Pt(2.5)
    p = nz.text_frame.paragraphs[0]
    p.text = "Zᵢ"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Yi: centro en (4.9, 3.75)
    ny = s9.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.35), Inches(3.2), Inches(1.1), Inches(1.1))
    ny.fill.solid()
    ny.fill.fore_color.rgb = GRAY_BG
    ny.line.color.rgb = OBS_COLOR
    ny.line.width = Pt(2.5)
    p = ny.text_frame.paragraphs[0]
    p.text = "Yᵢ"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = OBS_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Global Theta (outside plate): centro en (4.9, 1.0)
    ntheta = s9.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.35), Inches(0.7), Inches(1.1), Inches(0.7))
    ntheta.fill.solid()
    ntheta.fill.fore_color.rgb = WHITE
    ntheta.line.color.rgb = INK_COLOR
    ntheta.line.width = Pt(2)
    p = ntheta.text_frame.paragraphs[0]
    p.text = "θ"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR
    p.alignment = PP_ALIGN.CENTER

    # ARCOS VECTORIALES
    # 1. Xi -> Yi
    add_arrow(s9, Inches(2.75), Inches(2.7), Inches(4.35), Inches(3.4), dashed=False, color=CONTEXT_COLOR)
    # 2. Zi -> Yi
    add_arrow(s9, Inches(2.75), Inches(4.8), Inches(4.35), Inches(4.0), dashed=False, color=LATENT_COLOR)
    # 3. Xi -> Zi (discontinuo/punteado)
    add_arrow(s9, Inches(2.2), Inches(3.25), Inches(2.2), Inches(4.25), dashed=True, color=LINE_COLOR)
    # 4. Theta -> Yi
    add_arrow(s9, Inches(4.9), Inches(1.4), Inches(4.9), Inches(3.2), dashed=False, color=INK_COLOR)

    # Right side text
    tb9_r = s9.shapes.add_textbox(Inches(6.6), Inches(1.6), Inches(5.9), Inches(4.8))
    tf9_r = tb9_r.text_frame
    tf9_r.word_wrap = True

    p = tf9_r.paragraphs[0]
    p.text = "¿Qué se comparte y qué se repite?"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR

    p = tf9_r.add_paragraph()
    p.text = "• θ (Parámetros globales): Estructura de la red y CPTs fijas entrenadas. Fuera del plate porque se comparten en todo el análisis."
    p.font.size = Pt(15.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(8)

    p = tf9_r.add_paragraph()
    p.text = "• Xᵢ, Yᵢ, Zᵢ (Variables locales): Cada sitio o píxel i posee su propio valor de contexto, su medición instrumental y su condición inferida."
    p.font.size = Pt(15.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(8)

    p = tf9_r.add_paragraph()
    p.text = "• Supuesto de independencia condicional: Dados θ y X, los sitios son condicionalmente independientes en este modelo base."
    p.font.size = Pt(15.5)
    p.font.color.rgb = MUTED_COLOR
    p.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 10: Un modelo compartido para los píxeles de México (en español)
    # -------------------------------------------------------------
    s10 = add_base_slide("Un modelo compartido para los píxeles de México", 10,
        "Fuente: Descripción del procedimiento por el equipo (P016). "
        "Imagen: Mapa nacional de Integridad Ecosistémica con malla de píxeles en español. "
        "El modelo entrenado con parámetros θ se aplica de manera homogénea sobre los vectores de cada píxel nacional.")

    img_mapa = ROOT / "presentaciones/img/mapa_mexico_pixeles_iie_es.jpg"
    if not img_mapa.exists():
        img_mapa = ROOT / "presentaciones/img/mapa_mexico_pixeles_iie.jpg"
    if not img_mapa.exists():
        img_mapa = ROOT / "img/Mapa_México_Página_3.png"
    if img_mapa.exists():
        s10.shapes.add_picture(str(img_mapa), Inches(0.8), Inches(1.4), Inches(6.8), Inches(5.3))

    tb10_r = s10.shapes.add_textbox(Inches(7.9), Inches(1.5), Inches(4.6), Inches(5.1))
    tf10 = tb10_r.text_frame
    tf10.word_wrap = True

    p = tf10.paragraphs[0]
    p.text = "En cada píxel i"
    p.font.name = FONT_MAIN
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR

    p = tf10.add_paragraph()
    p.text = "Un vector reúne contexto Xᵢ (bioclima, suelo) y detección Yᵢ (sensores remotos)."
    p.font.size = Pt(15.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf10.add_paragraph()
    p.text = "El mismo modelo θ"
    p.font.name = FONT_MAIN
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(14)

    p = tf10.add_paragraph()
    p.text = "Los parámetros del modelo permanecen fijos. Cambian las entradas por píxel y se obtiene el diagnóstico de integridad nacional."
    p.font.size = Pt(15.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf10.add_paragraph()
    p.text = "Nota: La escala 0–1 del mapa refleja el índice sintético, no un porcentaje determinista ni una probabilidad escalar pura."
    p.font.size = Pt(13.5)
    p.font.color.rgb = MUTED_COLOR
    p.space_before = Pt(14)

    # -------------------------------------------------------------
    # SLIDE 11: Plates anidados (con arcos vectoriales)
    # -------------------------------------------------------------
    s11 = add_base_slide("Plates anidados: contexto regional y señales locales", 11,
        "Fuente: P016. "
        "El plate exterior representa clases o zonas biogeográficas j = 1,...,J (ej. Zonas de Holdridge). "
        "El plate interior representa píxeles i = 1,...,n_j pertenecientes a la zona j. "
        "Compartir contexto X_j no obliga a obtener el mismo IIE ni garantiza independencia espacial. "
        "El arco contextual hacia condición X_j -> Z_ij permanece en evaluación y se marca con trazo discontinuo como hipótesis abierta.")

    # Outer plate: Regions j
    p_out = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.2), Inches(6.0), Inches(4.5))
    p_out.fill.background()
    p_out.line.color.rgb = CONTEXT_COLOR
    p_out.line.width = Pt(2.5)

    tb_pout = s11.shapes.add_textbox(Inches(1.0), Inches(6.25), Inches(5.6), Inches(0.35))
    p = tb_pout.text_frame.paragraphs[0]
    p.text = "Zonas o clases biogeográficas   j = 1, ..., J"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR

    # Inner plate: Pixels i in j
    p_in = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.1), Inches(3.6), Inches(5.4), Inches(2.4))
    p_in.fill.background()
    p_in.line.color.rgb = MUTED_COLOR
    p_in.line.width = Pt(2)

    tb_pin = s11.shapes.add_textbox(Inches(1.3), Inches(5.6), Inches(5.0), Inches(0.35))
    p = tb_pin.text_frame.paragraphs[0]
    p.text = "Píxeles locales   i = 1, ..., nⱼ"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = MUTED_COLOR

    # Node: Xj (top inside outer plate) -> centro en (2.0, 2.9)
    nxj = s11.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.55), Inches(2.45), Inches(0.9), Inches(0.9))
    nxj.fill.solid()
    nxj.fill.fore_color.rgb = GRAY_BG
    nxj.line.color.rgb = CONTEXT_COLOR
    nxj.line.width = Pt(2)
    p = nxj.text_frame.paragraphs[0]
    p.text = "Xⱼ"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Nodes inside inner: Zij (centro en 2.4, 4.6), Yij (centro en 5.4, 4.6)
    nzij = s11.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.95), Inches(4.15), Inches(0.9), Inches(0.9))
    nzij.fill.solid()
    nzij.fill.fore_color.rgb = WHITE
    nzij.line.color.rgb = LATENT_COLOR
    nzij.line.width = Pt(2)
    p = nzij.text_frame.paragraphs[0]
    p.text = "Zᵢⱼ"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR
    p.alignment = PP_ALIGN.CENTER

    nyij = s11.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4.95), Inches(4.15), Inches(0.9), Inches(0.9))
    nyij.fill.solid()
    nyij.fill.fore_color.rgb = GRAY_BG
    nyij.line.color.rgb = OBS_COLOR
    nyij.line.width = Pt(2)
    p = nyij.text_frame.paragraphs[0]
    p.text = "Yᵢⱼ"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = OBS_COLOR
    p.alignment = PP_ALIGN.CENTER

    # Global theta: centro en (5.4, 1.45)
    nth = s11.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4.95), Inches(1.15), Inches(0.9), Inches(0.6))
    nth.fill.solid()
    nth.fill.fore_color.rgb = WHITE
    nth.line.color.rgb = INK_COLOR
    nth.line.width = Pt(2)
    p = nth.text_frame.paragraphs[0]
    p.text = "θ"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR
    p.alignment = PP_ALIGN.CENTER

    # ARCOS VECTORIALES
    # 1. Xj -> Zij (discontinuo/punteado)
    add_arrow(s11, Inches(2.0), Inches(3.35), Inches(2.4), Inches(4.15), dashed=True, color=LINE_COLOR)
    # 2. Xj -> Yij (sólido)
    add_arrow(s11, Inches(2.45), Inches(2.9), Inches(4.95), Inches(4.2), dashed=False, color=CONTEXT_COLOR)
    # 3. Zij -> Yij (sólido)
    add_arrow(s11, Inches(2.85), Inches(4.6), Inches(4.95), Inches(4.6), dashed=False, color=LATENT_COLOR)
    # 4. Theta -> Yij (sólido)
    add_arrow(s11, Inches(5.4), Inches(1.75), Inches(5.4), Inches(4.15), dashed=False, color=INK_COLOR)

    # Right side text
    tb11_r = s11.shapes.add_textbox(Inches(7.2), Inches(1.5), Inches(5.3), Inches(5.1))
    tf11 = tb11_r.text_frame
    tf11.word_wrap = True

    p = tf11.paragraphs[0]
    p.text = "Jerarquía espacial didáctica"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = INK_COLOR

    p = tf11.add_paragraph()
    p.text = "• Xⱼ (Contexto compartido): Zona bioclimática o clase que engloba muchos píxeles."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf11.add_paragraph()
    p.text = "• Yᵢⱼ (Detección local): Señales de satélite y cobertura del píxel i dentro de la zona j."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf11.add_paragraph()
    p.text = "• Zᵢⱼ (Condición inferida): Diagnóstico específico del píxel."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(6)

    p = tf11.add_paragraph()
    p.text = "Idea clave: Dos píxeles en la misma zona de Holdridge comparten Xⱼ pero pueden tener diferente diagnóstico si sus señales Yᵢⱼ difieren."
    p.font.size = Pt(14.5)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR
    p.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 12: Retroalimentación y tiempo: Redes Dinámicas (DBN en español)
    # -------------------------------------------------------------
    s12 = add_base_slide("Retroalimentación y tiempo: Redes Dinámicas (DBN)", 12,
        "Fuentes: Dean & Kanazawa (1989), Murphy (2002), Cafe-blog/posts/dynamic-bayesian-network/index.qmd. "
        "¿Cómo se representan los bucles de retroalimentación de la dinámica de sistemas sin violar la aciclicidad del DAG? "
        "Desdoblando el tiempo: el estado en t0 influye sobre el estado en t1 (Z_t0 -> Z_t1). "
        "Esto modela la memoria ecológica y la resiliencia como persistencia temporal.")

    tb12_top = s12.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(1.5))
    tf12 = tb12_top.text_frame
    tf12.word_wrap = True

    p = tf12.paragraphs[0]
    p.text = "Desdoblar los ciclos en el tiempo (t₀ → t₁)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR

    p = tf12.add_paragraph()
    p.text = "• La dinámica de sistemas clásica enfatiza bucles de retroalimentación (feedback loops). En un DAG estático no puede haber ciclos."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    p = tf12.add_paragraph()
    p.text = "• Solución natural: Redes Bayesianas Dinámicas (DBN). El estado en t₀ condiciona probabilísticamente la transición hacia t₁ [ Z(t₀) → Z(t₁) ]."
    p.font.size = Pt(14.5)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # Bottom Image: DBN Transition en español
    img_dbn = ROOT / "presentaciones/img/dbn_temporal_transition_es.jpg"
    if not img_dbn.exists():
        img_dbn = ROOT / "presentaciones/img/dbn_temporal_transition.jpg"
    if img_dbn.exists():
        s12.shapes.add_picture(str(img_dbn), Inches(0.8), Inches(3.0), Inches(11.733), Inches(3.8))

    # -------------------------------------------------------------
    # SLIDE 13: Supuestos clave y cautela interpretativa
    # -------------------------------------------------------------
    s13 = add_base_slide("Supuestos clave y cautela interpretativa", 13,
        "Revisión conceptual rigurosa: iie-teoria §16.2 y §21. "
        "1. Las flechas representan dependencias probabilísticas o hipótesis generativas, no causalidad mágica. "
        "2. Condicionar evidencia no es lo mismo que intervenir experimentalmente (do-calculus). "
        "3. El índice no es un porcentaje de masa física. "
        "4. La resolución espacial no añade información independiente automáticamente.")

    box13_w = Inches(5.6)
    box13_h = Inches(2.25)
    cards13 = [
        ("1. Arcos ≠ Causalidad demostrada",
         "Una flecha aprendida o trazada representa dependencia condicional. Las afirmaciones causales de intervención requieren supuestos experimentales adicionales.",
         CONTEXT_COLOR, Inches(0.8), Inches(1.55)),
        ("2. Condicionar ≠ Intervenir",
         "Observar que Y = y actualiza nuestra creencia diagnóstica (inferencia pasiva). Intervenir sobre Y mediante manejo [do(Y)] modifica la estructura del sistema.",
         OBS_COLOR, Inches(6.8), Inches(1.55)),
        ("3. Naturaleza de la probabilidad",
         "P(Integridad=alta) = 0.72 no significa que el ecosistema tenga 72% de materia viva. Expresa el grado de certidumbre diagnóstica dado el modelo.",
         LATENT_COLOR, Inches(0.8), Inches(4.2)),
        ("4. Escala y soporte de datos",
         "Mayor resolución de píxel no produce automáticamente observaciones independientes. El soporte espacial y los errores de medición deben considerarse.",
         ACCENT_GREEN, Inches(6.8), Inches(4.2)),
    ]

    for title_c, desc_c, col, px, py in cards13:
        rect = s13.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, px, py, box13_w, box13_h)
        rect.fill.solid()
        rect.fill.fore_color.rgb = WHITE
        rect.line.color.rgb = col
        rect.line.width = Pt(2)

        tf = rect.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.12)

        p = tf.paragraphs[0]
        p.text = title_c.rstrip(".")
        p.font.name = FONT_MAIN
        p.font.size = Pt(16.5)
        p.font.bold = True
        p.font.color.rgb = col

        p = tf.add_paragraph()
        p.text = desc_c
        p.font.name = FONT_MAIN
        p.font.size = Pt(13.5)
        p.font.color.rgb = INK_COLOR
        p.space_before = Pt(4)

    # -------------------------------------------------------------
    # SLIDE 14: Comprobación conceptual
    # -------------------------------------------------------------
    s14 = add_base_slide("Comprobación conceptual: preguntas para el grupo", 14,
        "Respuestas esperadas para el facilitador: "
        "Pregunta 1: Dos píxeles en la misma zona de vida comparten contexto X_j y usan el mismo modelo θ, "
        "pero pueden tener diferente diagnóstico Z_ij si sus señales observables Y_ij (NDVI, cobertura, biomasa) son distintas. "
        "Pregunta 2: Porque la salud es una propiedad sistémica de procesos e interacciones funcionales acopladas (como en la pecera), "
        "no una simple suma aritmética de piezas verdes aisladas.")

    qbox_w = Inches(11.7)
    qbox_h = Inches(2.25)

    # Question 1 Box
    q1 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), qbox_w, qbox_h)
    q1.fill.solid()
    q1.fill.fore_color.rgb = WHITE
    q1.line.color.rgb = CONTEXT_COLOR
    q1.line.width = Pt(2)

    tf_q1 = q1.text_frame
    tf_q1.word_wrap = True
    tf_q1.margin_left = tf_q1.margin_right = Inches(0.25)
    tf_q1.margin_top = Inches(0.15)

    p = tf_q1.paragraphs[0]
    p.text = "Pregunta 1 (Escala y diagnóstico)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(17.5)
    p.font.bold = True
    p.font.color.rgb = CONTEXT_COLOR

    p = tf_q1.add_paragraph()
    p.text = "Si dos píxeles de México pertenecen a la misma zona de vida de Holdridge (mismo contexto Xⱼ) y se evalúan con el mismo modelo θ, ¿por qué pueden obtener diagnósticos de integridad (Zᵢⱼ) completamente diferentes?"
    p.font.size = Pt(15)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # Question 2 Box
    q2 = s14.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.2), qbox_w, qbox_h)
    q2.fill.solid()
    q2.fill.fore_color.rgb = WHITE
    q2.line.color.rgb = LATENT_COLOR
    q2.line.width = Pt(2)

    tf_q2 = q2.text_frame
    tf_q2.word_wrap = True
    tf_q2.margin_left = tf_q2.margin_right = Inches(0.25)
    tf_q2.margin_top = Inches(0.15)

    p = tf_q2.paragraphs[0]
    p.text = "Pregunta 2 (Metáfora de salud e integridad)"
    p.font.name = FONT_MAIN
    p.font.size = Pt(17.5)
    p.font.bold = True
    p.font.color.rgb = LATENT_COLOR

    p = tf_q2.add_paragraph()
    p.text = "¿Por qué la metáfora de la pecera y la salud organísmica nos advierten que no podemos calcular la integridad ecosistémica simplemente sumando o promediando piezas verdes aisladas?"
    p.font.size = Pt(15)
    p.font.color.rgb = INK_COLOR
    p.space_before = Pt(4)

    # Guardar presentación
    try:
        prs.save(str(PPTX_PATH))
        print(f"Presentación generada exitosamente en: {PPTX_PATH}")
    except PermissionError:
        fallback_path = OUTPUT_DIR / "Introduccion-iie3t-redes-bayesianas-v4-es.pptx"
        prs.save(str(fallback_path))
        print(f"Nota: {PPTX_PATH.name} está abierto en PowerPoint. Se guardó copia en: {fallback_path}")


if __name__ == "__main__":
    create_deck()
